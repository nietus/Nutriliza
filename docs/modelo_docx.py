"""Monta documentos .docx com o mesmo template visual do trabalho Escaliza (PUC Minas).

O template (docs/template_puc.docx) é aberto, a capa e a folha de rosto são
reaproveitadas com os textos trocados e o corpo é reconstruído clonando a
formatação direta dos parágrafos-modelo: título, seção, subseção, corpo,
lista, legenda de tabela, tabela e fonte.
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Emu

TEMPLATE = Path(__file__).resolve().parent / "template_puc.docx"

# índices dos elementos-modelo no corpo do template
_CAPA_FIM = 31          # parágrafo com a quebra de seção da capa
_P_INTEGRANTES, _P_TITULO, _P_H1, _P_CORPO = 32, 33, 34, 35
_P_H2, _P_LEGENDA, _TBL, _P_FONTE, _P_LISTA = 36, 37, 38, 39, 40


def _textos(el):
    return el.iter(qn("w:t"))


def _trocar_texto(par, novo: str) -> None:
    ts = list(_textos(par))
    ts[0].text = novo
    for t in ts[1:]:
        t.text = ""


class Documento:
    def __init__(self, titulo: str, subtitulo: str, integrantes: list[str], disciplina: str,
                 professor: str, descricao: str, cidade_ano=("Belo Horizonte – MG", "2026")):
        self.doc = docx.Document(str(TEMPLATE))
        corpo = self.doc.element.body
        kids = list(corpo)
        self.proto = {k: copy.deepcopy(kids[i]) for k, i in {
            "integrantes": _P_INTEGRANTES, "titulo": _P_TITULO, "h1": _P_H1, "corpo": _P_CORPO,
            "h2": _P_H2, "legenda": _P_LEGENDA, "tabela": _TBL, "fonte": _P_FONTE,
            "lista": _P_LISTA}.items()}
        self.proto["capa_figura"] = copy.deepcopy(kids[0])
        self._capa(kids, titulo, subtitulo, integrantes, disciplina, professor, descricao, cidade_ano)
        for el in kids[_CAPA_FIM + 1:]:
            if el.tag != qn("w:sectPr"):
                corpo.remove(el)
        self._sectpr = corpo[-1]
        self._cabecalho_rodape(disciplina, integrantes)
        self.n_tab = self.n_fig = self.n_quadro = 0
        self._integrantes_e_titulo(integrantes, titulo, subtitulo)

    # ---------- capa ----------
    def _capa(self, kids, titulo, subtitulo, integrantes, disciplina, professor, descricao, cidade_ano):
        _trocar_texto(kids[3], f"Disciplina: {disciplina}")
        for bloco_ini, bloco_fim, t_titulo, t_sub in ((4, 11, 12, 13), (17, 24, 25, 26)):
            nomes = kids[bloco_ini:bloco_fim + 1]
            for k, par in enumerate(nomes):
                if k < len(integrantes):
                    _trocar_texto(par, integrantes[k])
                else:
                    par.getparent().remove(par)
            _trocar_texto(kids[t_titulo], titulo.upper())
            _trocar_texto(kids[t_sub], subtitulo)
        _trocar_texto(kids[27], descricao)
        _trocar_texto(kids[28], f"Professor: {professor}")
        for i in (14, 29):
            _trocar_texto(kids[i], cidade_ano[0])
        for i in (15, 30):
            _trocar_texto(kids[i], cidade_ano[1])

    def _cabecalho_rodape(self, disciplina, integrantes):
        sec = self.doc.sections[-1]
        for parte in (sec.header, sec.footer):
            for t in parte._element.iter(qn("w:t")):
                if t.text == "Tópicos III":
                    t.text = disciplina
                elif t.text == "Integrantes":
                    t.text = ", ".join(integrantes)

    def _integrantes_e_titulo(self, integrantes, titulo, subtitulo):
        p = copy.deepcopy(self.proto["integrantes"])
        ts = list(_textos(p))
        ts[0].text = "Integrantes: "
        ts[1].text = ", ".join(integrantes)
        self._add(p)
        self._add(self._par("titulo", f"{titulo} – {subtitulo}"))

    # ---------- utilitários ----------
    def _add(self, el):
        self._sectpr.addprevious(el)
        return el

    def _par(self, tipo: str, texto: str):
        p = copy.deepcopy(self.proto[tipo])
        runs = p.findall(qn("w:r"))
        for r in runs[1:]:
            p.remove(r)
        list(_textos(runs[0]))[0].text = texto
        return p

    def _par_rico(self, tipo: str, texto: str):
        """Parágrafo com **negrito** e *itálico* simples."""
        p = copy.deepcopy(self.proto[tipo])
        runs = p.findall(qn("w:r"))
        modelo = copy.deepcopy(runs[-1])
        for r in runs:
            p.remove(r)
        rpr = modelo.find(qn("w:rPr"))
        for b in rpr.findall(qn("w:b")):
            rpr.remove(b)
        for parte in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
            if not parte:
                continue
            r = copy.deepcopy(modelo)
            r_pr = r.find(qn("w:rPr"))
            if parte.startswith("**"):
                r_pr.insert(0, r_pr.makeelement(qn("w:b"), {}))
                parte = parte[2:-2]
            elif parte.startswith("*"):
                r_pr.insert(0, r_pr.makeelement(qn("w:i"), {}))
                parte = parte[1:-1]
            t = r.find(qn("w:t"))
            t.text = parte
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            p.append(r)
        return p

    # ---------- API ----------
    def secao(self, texto):
        self._add(self._par("h1", texto))

    def subsecao(self, texto):
        self._add(self._par("h2", texto))

    def paragrafo(self, texto):
        self._add(self._par_rico("corpo", texto))

    def equacao(self, texto):
        p = self._par("corpo", texto)
        jc = p.find(qn("w:pPr")).find(qn("w:jc"))
        jc.set(qn("w:val"), "center")
        r_pr = p.find(qn("w:r")).find(qn("w:rPr"))
        r_pr.insert(0, r_pr.makeelement(qn("w:i"), {}))
        self._add(p)

    def lista(self, itens: list[str]):
        for item in itens:
            self._add(self._par_rico("lista", item))

    def fonte(self, texto):
        self._add(self._par("fonte", texto))

    def tabela(self, legenda: str, cabecalho: list[str], linhas: list[list], larguras: list[float],
               fonte: str | None = "Fonte: elaboração do grupo.", tipo: str = "Tabela",
               alinhar_direita: set[int] | None = None):
        if tipo == "Quadro":
            self.n_quadro += 1
            n = self.n_quadro
        else:
            self.n_tab += 1
            n = self.n_tab
        self._add(self._par("legenda", f"{tipo} {n} – {legenda}"))
        tbl = copy.deepcopy(self.proto["tabela"])
        total = 9638
        cols = [int(total * w / sum(larguras)) for w in larguras]
        cols[-1] = total - sum(cols[:-1])
        grid = tbl.find(qn("w:tblGrid"))
        for g in list(grid):
            grid.remove(g)
        for w in cols:
            grid.append(grid.makeelement(qn("w:gridCol"), {qn("w:w"): str(w)}))
        trs = tbl.findall(qn("w:tr"))
        tr_cab, tr_dado = trs[0], trs[1]
        for tr in trs:
            tbl.remove(tr)

        def linha(modelo, valores, cabeca=False):
            tr = copy.deepcopy(modelo)
            tcs = tr.findall(qn("w:tc"))
            tc_negrito, tc_normal = tcs[0], tcs[-1]
            for tc in tcs:
                tr.remove(tc)
            for k, (v, w) in enumerate(zip(valores, cols)):
                tc = copy.deepcopy(tc_negrito if (cabeca or k == 0) else tc_normal)
                tc.find(qn("w:tcPr")).find(qn("w:tcW")).set(qn("w:w"), str(w))
                par = tc.find(qn("w:p"))
                list(_textos(par))[0].text = str(v)
                if alinhar_direita and k in alinhar_direita:
                    par.find(qn("w:pPr")).find(qn("w:jc")).set(qn("w:val"), "right")
                tr.append(tc)
            tbl.append(tr)

        linha(tr_cab, cabecalho, cabeca=True)
        for valores in linhas:
            linha(tr_dado, valores)
        # mantém a tabela inteira na mesma página (todas as linhas menos a última "seguem a próxima")
        for tr in tbl.findall(qn("w:tr"))[:-1]:
            for ppr in tr.iter(qn("w:pPr")):
                if ppr.find(qn("w:keepNext")) is None:
                    ppr.insert(0, ppr.makeelement(qn("w:keepNext"), {}))
        self._add(tbl)
        if fonte:
            self.fonte(fonte)

    def figura(self, caminho: Path, legenda: str, largura_cm: float = 16.0,
               fonte: str | None = "Fonte: elaboração do grupo."):
        self.n_fig += 1
        self._add(self._par("legenda", f"Figura {self.n_fig} – {legenda}"))
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(str(caminho), width=Emu(int(largura_cm * 360000)))
        self._add(par._p)
        if fonte:
            self.fonte(fonte)

    def referencia(self, texto):
        p = self._par_rico("corpo", texto)
        p.find(qn("w:pPr")).find(qn("w:jc")).set(qn("w:val"), "left")
        self._add(p)

    def salvar(self, caminho: Path):
        self.doc.save(str(caminho))
